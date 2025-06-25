# apps/sez/views/export_products.py
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from drf_spectacular.utils import extend_schema, OpenApiResponse
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from apps.sez.permissions import ClearanceInvoiceItemsPermission
from apps.shtrih.models import Products


@extend_schema(tags=['Clearance Workflow'])
@extend_schema(
    summary="Export Products for ClearanceInvoice",
    description="Generate and download XLSX with products linked to the specified ClearanceInvoice.",
    responses={
        200: OpenApiResponse(
            description="XLSX file",
        )
    },
)
class ClearanceInvoiceProductsExportView(APIView):
    """
    Export Products for a given ClearanceInvoice as an XLSX file.

    Excel columns:
    1) Product ID
    2) Model Short Name
    3) Barcode
    """
    permission_classes = [IsAuthenticated, ClearanceInvoiceItemsPermission]

    def get(self, request, invoice_id: int) -> HttpResponse:
        # Fetch products related to the given invoice
        products_qs = (
            Products.objects
            .filter(cleared_id=invoice_id)
            .select_related('model__name')
            .order_by('id')
        )

        # Create workbook and sheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Invoice_{invoice_id}_Products"

        # Write header
        headers = ["Product ID", "Model Short Name", "Barcode"]
        ws.append(headers)

        # Write data rows
        for product in products_qs:
            product_id = product.id
            model_short = product.model.name.short_name if product.model and product.model.name else ''
            barcode = product.barcode or ''
            ws.append([product_id, model_short, barcode])

        # Adjust column widths and alignment
        # Set wider columns
        column_widths = {
            1: 15,  # Product ID
            2: 20,  # Model Short Name
            3: 30,  # Barcode
        }
        for col_idx, width in column_widths.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        # Align all cells to left
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.alignment = Alignment(horizontal='left')

        # Save workbook to in-memory buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        # Prepare response
        filename = f"clearance_invoice_{invoice_id}_products.xlsx"
        response = HttpResponse(
            buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
